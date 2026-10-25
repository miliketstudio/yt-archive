import json
import os
import shutil
import subprocess
import time
import zipfile
from pathlib import Path

import numpy as np

WORK = Path("work")
FRAMES = WORK / "frames"
OUT = Path("out")

JOB_ID = os.environ["JOB_ID"]
URL = os.environ["URL"]
MODE = os.environ.get("MODE", "both")  # slides | video | both
THRESHOLD = float(os.environ.get("SENSITIVITY") or "3.0")
MAX_FRAMES = int(os.environ.get("MAX_FRAMES") or "300")
REPO = os.environ.get("REPO", "")
TAG = f"job-{JOB_ID}"
MAX_ASSET_BYTES = 2 * 1024**3 - 1

db = None


def init_db():
    global db
    raw = os.environ.get("FIREBASE_SERVICE_ACCOUNT")
    if not raw:
        print("No FIREBASE_SERVICE_ACCOUNT set, skipping Firestore updates")
        return
    import firebase_admin
    from firebase_admin import credentials, firestore

    firebase_admin.initialize_app(credentials.Certificate(json.loads(raw)))
    db = firestore.client()


def update(**fields):
    print("job update:", fields)
    if db:
        db.collection("jobs").document(JOB_ID).set(fields, merge=True)


def run(cmd, check=True, **kw):
    print("+", " ".join(str(c) for c in cmd))
    return subprocess.run(cmd, check=check, **kw)


def ytdlp_base():
    cmd = ["yt-dlp", "--no-playlist", "--js-runtimes", "node"]
    cookies = os.environ.get("YT_COOKIES")
    if cookies:
        Path("cookies.txt").write_text(cookies)
        cmd += ["--cookies", "cookies.txt"]
    return cmd


def download():
    meta = run(
        ytdlp_base() + ["--dump-single-json", URL],
        stdout=subprocess.PIPE,
        text=True,
    )
    info = json.loads(meta.stdout)
    fmt = (
        "bv*[vcodec^=avc1][height<=1080]+ba[acodec^=mp4a]"
        "/bv*[height<=1080]+ba/b[height<=1080]/b"
    )
    run(
        ytdlp_base()
        + [
            "-f", fmt,
            "--merge-output-format", "mp4",
            "-o", str(WORK / "video.%(ext)s"),
            URL,
        ]
    )
    return info, WORK / "video.mp4"


def find_changes(video):
    """Sample 1 frame/second at low resolution and return the second marks
    where the screen changed and then settled."""
    w, h = 160, 90
    size = w * h
    cmd = [
        "ffmpeg", "-v", "error", "-i", str(video),
        "-vf", f"fps=1,scale={w}:{h},format=gray",
        "-f", "rawvideo", "-",
    ]
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE)
    kept, last_kept, prev, pending, t = [], None, None, 0, 0
    while True:
        buf = proc.stdout.read(size)
        if len(buf) < size:
            break
        cur = np.frombuffer(buf, dtype=np.uint8).astype(np.int16)
        if last_kept is None:
            kept.append(t)
            last_kept = cur
        else:
            changed = np.abs(cur - last_kept).mean() > THRESHOLD
            settled = np.abs(cur - prev).mean() < 1.0
            pending = pending + 1 if changed else 0
            # keep once the picture stops moving; give up waiting after 8s
            if changed and (settled or pending >= 8):
                kept.append(t)
                last_kept = cur
                pending = 0
        prev = cur
        t += 1
    proc.wait()
    if len(kept) > MAX_FRAMES:
        step = len(kept) / MAX_FRAMES
        kept = [kept[int(i * step)] for i in range(MAX_FRAMES)]
    return kept


def extract(video, times):
    FRAMES.mkdir(parents=True, exist_ok=True)
    files = []
    for i, t in enumerate(times, 1):
        f = FRAMES / f"{i:04d}.jpg"
        subprocess.run(
            ["ffmpeg", "-v", "error", "-y", "-ss", str(t), "-i", str(video),
             "-frames:v", "1", "-q:v", "2", str(f)],
            check=True,
        )
        if f.exists():
            files.append(f)
    return files


def build_pdf(files, out):
    import img2pdf

    out.write_bytes(img2pdf.convert([str(f) for f in files]))


def build_pptx(files, times, out):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    w, h = Image.open(files[0]).size
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = int(prs.slide_width * h / w)
    blank = prs.slide_layouts[6]
    for f, t in zip(files, times):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(f), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
        slide.notes_slide.notes_text_frame.text = time.strftime(
            "%H:%M:%S", time.gmtime(t)
        )
    prs.save(str(out))


def build_zip(files, times, out):
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        for f in files:
            z.write(f, f.name)
        manifest = [{"file": f.name, "t": t} for f, t in zip(files, times)]
        z.writestr("frames.json", json.dumps(manifest))


def publish_release(title, assets):
    run(["gh", "release", "delete", TAG, "--cleanup-tag", "-y", "-R", REPO],
        check=False)
    run(
        ["gh", "release", "create", TAG, "-R", REPO,
         "--title", title[:200], "--notes", URL]
        + [str(a) for a in assets]
    )


def main():
    init_db()
    OUT.mkdir(exist_ok=True)
    WORK.mkdir(exist_ok=True)
    update(status="running", url=URL, mode=MODE, releaseTag=TAG)
    try:
        info, video = download()
        title = info.get("title") or "video"
        update(title=title, duration=info.get("duration"),
               thumbnail=info.get("thumbnail"))

        assets = []
        if MODE in ("slides", "both"):
            times = find_changes(video)
            files = extract(video, times)
            times = times[: len(files)]
            if not files:
                raise RuntimeError("No frames could be extracted")
            build_pdf(files, OUT / "slides.pdf")
            build_pptx(files, times, OUT / "slides.pptx")
            build_zip(files, times, OUT / "frames.zip")
            assets += [OUT / "slides.pdf", OUT / "slides.pptx", OUT / "frames.zip"]
            update(frameCount=len(files))

        if MODE in ("video", "both"):
            if video.stat().st_size > MAX_ASSET_BYTES:
                update(videoSkipped="Video is over GitHub's 2 GB file limit")
            else:
                final = OUT / "video.mp4"
                shutil.copy(video, final)
                assets.append(final)

        for a in list(assets):
            if a.stat().st_size > MAX_ASSET_BYTES:
                assets.remove(a)

        publish_release(title, assets)
        update(status="done", assets={a.name: a.stat().st_size for a in assets})
    except Exception as e:
        update(status="error", error=str(e)[:500])
        raise


if __name__ == "__main__":
    main()
